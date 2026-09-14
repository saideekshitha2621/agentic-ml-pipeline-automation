"""Reports screen: CSV / Excel (openpyxl) / PDF (reportlab) export."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_csv(clustered_df: pd.DataFrame, out_path: Path) -> Path:
    clustered_df.to_csv(out_path, index=False)
    return out_path


def _excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    """openpyxl can't write dict/list cell values (e.g. the `params` column) — stringify them."""
    safe = df.copy()
    for col in safe.columns:
        if safe[col].map(lambda v: isinstance(v, (dict, list))).any():
            safe[col] = safe[col].map(lambda v: str(v) if isinstance(v, (dict, list)) else v)
    return safe


def export_excel(
    leaderboard_df: pd.DataFrame,
    clustered_df: pd.DataFrame,
    interpretation: dict,
    out_path: Path,
) -> Path:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Leaderboard"
    for row in dataframe_to_rows(_excel_safe(leaderboard_df), index=False, header=True):
        ws1.append(row)
    for cell in ws1[1]:
        cell.font = Font(bold=True)

    ws2 = wb.create_sheet("Clustered Data")
    for row in dataframe_to_rows(clustered_df, index=False, header=True):
        ws2.append(row)
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    ws3 = wb.create_sheet("Cluster Profiles")
    profiles_df = pd.DataFrame(interpretation.get("profiles", []))
    if not profiles_df.empty:
        for row in dataframe_to_rows(profiles_df, index=False, header=True):
            ws3.append(row)
        for cell in ws3[1]:
            cell.font = Font(bold=True)

    ws4 = wb.create_sheet("Cluster Summaries")
    ws4.append(["Cluster", "Suggested Name", "Business Summary"])
    for cell in ws4[1]:
        cell.font = Font(bold=True)
    names = interpretation.get("suggested_names", {})
    summaries = interpretation.get("summaries", {})
    for cluster_id, summary in summaries.items():
        ws4.append([cluster_id, names.get(cluster_id, ""), summary])

    wb.save(out_path)
    return out_path


def export_pdf(
    chosen_model_info: dict,
    leaderboard_df: pd.DataFrame,
    interpretation: dict,
    out_path: Path,
) -> Path:
    doc = SimpleDocTemplate(str(out_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("Unsupervised Clustering — Evaluation Report", styles["Title"]), Spacer(1, 16)]

    story.append(Paragraph("Chosen Model", styles["Heading2"]))
    model_lines = [
        f"Algorithm: {chosen_model_info.get('algorithm')}",
        f"Parameters: {chosen_model_info.get('params')}",
        f"Clusters: {chosen_model_info.get('n_clusters')}  |  Noise points: {chosen_model_info.get('n_noise')}",
        f"Silhouette: {chosen_model_info.get('silhouette_score')}  |  "
        f"Davies-Bouldin: {chosen_model_info.get('davies_bouldin_score')}  |  "
        f"Calinski-Harabasz: {chosen_model_info.get('calinski_harabasz_score')}",
    ]
    for line in model_lines:
        story.append(Paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 16))

    story.append(Paragraph("Leaderboard (Top 10)", styles["Heading2"]))
    top = leaderboard_df.head(10)
    table_data = [["Algorithm", "Params", "Sil.", "DB", "CH", "Rank"]] + [
        [
            r["algorithm"],
            str(r["params"])[:40],
            r.get("silhouette_score"),
            r.get("davies_bouldin_score"),
            r.get("calinski_harabasz_score"),
            r.get("rank"),
        ]
        for r in top.to_dict("records")
    ]
    table = Table(table_data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4C72B0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 16))

    story.append(Paragraph("Cluster Summaries", styles["Heading2"]))
    names = interpretation.get("suggested_names", {})
    for cluster_id, summary in interpretation.get("summaries", {}).items():
        name = names.get(cluster_id, cluster_id)
        story.append(Paragraph(f"<b>Cluster {cluster_id} — {name}</b>", styles["Normal"]))
        story.append(Paragraph(summary, styles["Normal"]))
        story.append(Spacer(1, 8))

    doc.build(story)
    return out_path
